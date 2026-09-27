# Trích xuất thông tin từ folder cover_letters ra file Excel

# ===== IMPORT =====
import os
import re
from docx import Document
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ===== INPUT =====
folder_path = "cover_letters"
output_file = "cover_letters_result.xlsx"


# ===== FUNCTIONS =====
def docx_to_text(file_path):
    doc = Document(file_path)

    text = ""

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text.strip() + "\n"

    return text


def tim_thong_tin(pattern, text):
    result = re.search(
        pattern,
        text,
        re.IGNORECASE | re.MULTILINE
    )

    if result:
        value = result.group(1).strip()

        value = re.sub(r"\s+", " ", value)

        if value:
            return value

    return "Không tìm thấy"


# ===== TẠO FILE EXCEL =====
workbook = Workbook()
sheet = workbook.active
sheet.title = "Cover Letters"


# ===== HEADER =====
headers = [
    "Tên file",
    "Họ và tên",
    "Giới tính",
    "Ngày sinh",
    "Nơi sinh",
    "Nguyên quán",
    "Hộ khẩu thường trú",
    "Chỗ ở hiện nay",
    "Điện thoại",
    "Dân tộc",
    "Tôn giáo",
    "CCCD/CMND",
    "Ngày cấp",
    "Nơi cấp",
    "Trình độ văn hóa",
    "Sở trường"
]

sheet.append(headers)


# ===== BIẾN ĐẾM =====
success_count = 0
error_count = 0


# ===== PROCESS =====
for file_name in os.listdir(folder_path):

    if file_name.lower().endswith(".docx"):

        file_path = os.path.join(
            folder_path,
            file_name
        )

        try:
            text = docx_to_text(file_path)

            ho_ten = tim_thong_tin(
                r"Họ và tên:\s*(.*?)\s+Nam/Nữ:",
                text
            )

            gioi_tinh = tim_thong_tin(
                r"Nam/Nữ:\s*([^\n]+)",
                text
            )

            ngay_sinh = tim_thong_tin(
                r"Sinh ngày:\s*(.*?)\s+Nơi sinh:",
                text
            )

            noi_sinh = tim_thong_tin(
                r"Nơi sinh:\s*([^\n]+)",
                text
            )

            nguyen_quan = tim_thong_tin(
                r"Nguyên quán:\s*([^\n]+)",
                text
            )

            ho_khau = tim_thong_tin(
                r"Nơi đăng ký hộ khẩu thường trú:\s*([^\n]+)",
                text
            )

            cho_o = tim_thong_tin(
                r"Chỗ ở hiện nay:\s*([^\n]+)",
                text
            )

            dien_thoai = tim_thong_tin(
                r"Điện thoại liên hệ:\s*([^\n]+)",
                text
            )

            dan_toc = tim_thong_tin(
                r"Dân tộc:\s*(.*?)\s+Tôn giáo",
                text
            )

            ton_giao = tim_thong_tin(
                r"Tôn giáo:\s*([^\n]+)",
                text
            )

            cccd = tim_thong_tin(
                r"Số CCCD/CMND:\s*(\d+)",
                text
            )

            ngay_cap = tim_thong_tin(
                r"Cấp ngày:\s*([^\s]+)",
                text
            )

            noi_cap = tim_thong_tin(
                r"Nơi cấp:\s*([^\n]+)",
                text
            )

            trinh_do = tim_thong_tin(
                r"Trình độ văn hóa:\s*([^\n]+)",
                text
            )

            so_truong = tim_thong_tin(
                r"Sở trường:\s*([^\n]+)",
                text
            )


            # ===== OUTPUT TO EXCEL =====
            sheet.append([
                file_name,
                ho_ten,
                gioi_tinh,
                ngay_sinh,
                noi_sinh,
                nguyen_quan,
                ho_khau,
                cho_o,
                dien_thoai,
                dan_toc,
                ton_giao,
                cccd,
                ngay_cap,
                noi_cap,
                trinh_do,
                so_truong
            ])

            success_count += 1

        except Exception as error:

            error_count += 1

            print(
                "Lỗi file:",
                file_name,
                "-",
                error
            )


# ===== TỰ ĐỘNG ĐIỀU CHỈNH ĐỘ RỘNG CỘT =====
for column_cells in sheet.columns:

    max_length = 0

    column_letter = get_column_letter(
        column_cells[0].column
    )

    for cell in column_cells:

        if cell.value is not None:

            cell_length = len(
                str(cell.value)
            )

            if cell_length > max_length:
                max_length = cell_length

    sheet.column_dimensions[
        column_letter
    ].width = max_length + 2


# ===== LƯU FILE =====
workbook.save(output_file)


# ===== OUTPUT =====
print("\n===== KẾT QUẢ =====")

print(
    "Số file đã xử lý thành công:",
    success_count
)

print(
    "Số file bị lỗi:",
    error_count
)

print(
    "Đã xuất dữ liệu ra file:",
    output_file
)