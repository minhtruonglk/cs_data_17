# Bài: Quản lý thông tin sinh viên và lưu vào Excel

# ===== IMPORT =====
import os
import re
from datetime import datetime, date
from openpyxl import Workbook, load_workbook


file_name = "students.xlsx"


# ===== FUNCTIONS =====
def kiem_tra_email(email):
    pattern = r"^[\w\.-]+@[\w\.-]+\.\w+$"
    return re.match(pattern, email) is not None


def kiem_tra_so_dien_thoai(phone):
    return phone.isdigit() and len(phone) == 10 and phone.startswith("0")


def nhap_ngay_sinh():
    while True:
        ngay_sinh_text = input("Nhập ngày sinh (dd/mm/yyyy): ")

        try:
            ngay_sinh = datetime.strptime(
                ngay_sinh_text,
                "%d/%m/%Y"
            ).date()

            if ngay_sinh > date.today():
                print("Ngày sinh không hợp lệ")
                continue

            return ngay_sinh

        except ValueError:
            print("Ngày sinh không hợp lệ. Vui lòng nhập lại")


def tinh_tuoi(ngay_sinh):
    hom_nay = date.today()

    tuoi = hom_nay.year - ngay_sinh.year

    if (hom_nay.month, hom_nay.day) < (ngay_sinh.month, ngay_sinh.day):
        tuoi -= 1

    return tuoi


# ===== MỞ HOẶC TẠO FILE EXCEL =====
if os.path.exists(file_name):
    workbook = load_workbook(file_name)
    sheet = workbook.active

else:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    sheet.append([
        "Mã sinh viên",
        "Họ tên",
        "Lớp",
        "Email",
        "Số điện thoại",
        "Ngày sinh",
        "Tuổi"
    ])


# ===== INPUT =====
while True:
    print("\n===== NHẬP THÔNG TIN SINH VIÊN =====")

    ma_sinh_vien = input("Nhập mã sinh viên: ")
    ho_ten = input("Nhập họ tên: ")
    lop = input("Nhập lớp: ")


    while True:
        email = input("Nhập email: ")

        if kiem_tra_email(email):
            break

        print("Email sai định dạng. Vui lòng nhập lại")


    while True:
        so_dien_thoai = input("Nhập số điện thoại: ")

        if kiem_tra_so_dien_thoai(so_dien_thoai):
            break

        print("Số điện thoại không hợp lệ. Vui lòng nhập lại")


    ngay_sinh = nhap_ngay_sinh()


    # ===== PROCESS =====
    tuoi = tinh_tuoi(ngay_sinh)


    # ===== OUTPUT =====
    sheet.append([
        ma_sinh_vien,
        ho_ten,
        lop,
        email,
        so_dien_thoai,
        ngay_sinh.strftime("%d/%m/%Y"),
        tuoi
    ])

    workbook.save(file_name)

    print("Đã lưu sinh viên vào", file_name)


    # ===== NHẬP TIẾP =====
    while True:
        tiep_tuc = input("Bạn có muốn nhập tiếp không? (Y/N): ").upper()

        if tiep_tuc == "Y" or tiep_tuc == "N":
            break

        print("Vui lòng chỉ nhập Y hoặc N")

    if tiep_tuc == "N":
        break


print("\nKết thúc chương trình")